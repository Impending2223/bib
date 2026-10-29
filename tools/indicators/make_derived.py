"""Real interest rates and the stance of policy: series computed from the indicators' own files, and one outside
estimate (the owner's request, Oct. 28, 2026: "Rising nominal rates are hot ... but rising real rates are cold";
"a whole-of-economy measure"; "What was the standard then ... What's the current (i.e., retrospective) standard?").

    python3 -I tools/indicators/make_derived.py --hlw PATH

PATH: the New York Fed's Holston-Laubach-Williams workbook (https://www.newyorkfed.org/research/policy/rstar, "HLW
current estimates", the .xlsx), read with openpyxl; without --hlw the r* row is kept as it stands. The build needs
neither. Writes, in indicators/:

    real-bill-cpi, real-10-cpi    monthly: the bill (the 10-year) less the CPI's rise over the year, each vintage's
    real-bill, real-10            quarterly: the quarter's average bill (10-year) less the GNP deflator's rise over the
                                  year: as first reported, the mean of the four quarters' annualized changes as each
                                  was released (a ratio across the base-year change of Aug. 1965 would mix 1954 and
                                  1958 prices); today, the deflator against its level a year before
    taylor-gap                    quarterly: the quarter's average federal funds rate less Taylor's rule (1993),
                                  2 + inflation + 1/2 (inflation - 2) + 1/2 (output gap): then, inflation as above and
                                  the gap by Okun's law as the Council used it, 3 percent of GNP for each point of
                                  unemployment (as first reported, the quarter's average) above 4; today, the
                                  deflator's rise and the CBO's gap (gap-cbo)
    rstar-gap                     quarterly, today only: HLW's real rate (their policy rate less their expected
                                  inflation) less their r* (US, one-sided)
    spread                        monthly, today only: the 10-year less the bill

Rules: a derived figure exists only where every input does; it rounds to two decimals; its change, from the period
before, in points. As first reported, a figure is "released" on the latest release among its inputs (the market
rates being known at the month's end, the CPI's or the deflator's release; for the Taylor row, the later of the
deflator's and the quarter's last unemployment release). The statements of concept are in indicators/concepts.yaml.
"""
import argparse
import datetime
import os
import sys

import yaml

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DIR = os.path.join(ROOT, "indicators")


def load(sid):
    with open(os.path.join(DIR, sid + ".yaml"), encoding="utf-8") as f:
        return yaml.safe_load(f)


def by_p(s):
    return {str(r["p"]): r for r in s["rows"]}


def quarter_of(p):
    y, m = int(p[:4]), int(p[5:7])
    return f"{y}Q{(m - 1) // 3 + 1}"


def months_of(q):
    y, n = int(q[:4]), int(q[-1])
    return [f"{y}-{m:02d}" for m in range(3 * n - 2, 3 * n + 1)]


def prev_q(q, k=1):
    y, n = int(q[:4]), int(q[-1])
    i = y * 4 + n - 1 - k
    return f"{i // 4}Q{i % 4 + 1}"


def prev_m(p, k=1):
    y, m = int(p[:4]), int(p[5:7])
    i = y * 12 + m - 1 - k
    return f"{i // 12}-{i % 12 + 1:02d}"


def qavg(s, q, field="now"):
    vals = [s.get(m, {}).get(field) for m in months_of(q)]
    return None if any(v is None for v in vals) else sum(vals) / 3


def latest(*dates):
    ds = [str(d) for d in dates if d]
    return max(ds) if ds else None


def r2(v):
    return None if v is None else round(v + 0.0, 2)


def with_changes(rows, freq):
    prev = prev_m if freq == "M" else prev_q
    idx = {r["p"]: r for r in rows}
    for r in rows:
        b = idx.get(prev(r["p"]))
        for f, c in (("first", "chg"), ("now", "chg_now")):
            if r.get(f) is not None and b is not None and b.get(f) is not None:
                r[c] = r2(r[f] - b[f])
    return rows


def series(sid, name, freq, then, now, rows):
    out = {"id": sid, "name": name, "freq": freq, "kind": "derived", "change": "pts", "then": then, "now": now,
           "rows": with_changes([r for r in rows if r.get("first") is not None or r.get("now") is not None], freq)}
    for r in out["rows"]:
        for k in [k for k, v in r.items() if v is None]:
            del r[k]
    return out


def write(s):
    path = os.path.join(DIR, s["id"] + ".yaml")
    with open(path, "w", encoding="utf-8") as f:
        f.write(f"# Derived by tools/indicators/make_derived.py; never edit. Statement: indicators/concepts.yaml.\n")
        yaml.safe_dump(s, f, allow_unicode=True, sort_keys=False, width=1000)
    print(f"indicators/{s['id']}.yaml: {len(s['rows'])} rows")


def hlw_rows(path):
    """The workbook's US r* (sheet "HLW Estimates", the third block) and its inputs ("US input data": the policy
    rate and the expected inflation), by quarter."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rstar, inp = {}, {}
    for row in wb["HLW Estimates"].iter_rows(min_row=7, values_only=True):
        if isinstance(row[0], datetime.datetime) and isinstance(row[10], (int, float)):
            rstar[f"{row[0].year}Q{(row[0].month - 1) // 3 + 1}"] = row[10]
    for row in wb["US input data"].iter_rows(min_row=2, values_only=True):
        if isinstance(row[0], datetime.datetime):
            inp[f"{row[0].year}Q{(row[0].month - 1) // 3 + 1}"] = (row[3], row[4])
    final = wb["HLW Estimates"].cell(row=3, column=1).value
    return rstar, inp, final


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--hlw", help="the New York Fed's HLW workbook (.xlsx)")
    a = ap.parse_args()
    bill, ten, ff = by_p(load("treasury-bill")), by_p(load("treasury-10-year")), by_p(load("federal-funds"))
    cpi, defl, ue, gap = by_p(load("cpi")), by_p(load("deflator")), by_p(load("unemployment")), by_p(load("gap-cbo"))
    months = sorted(p for p in bill if p in ten)
    quarters = sorted(p for p in defl if "Q" in p)

    # monthly, the CPI
    cpi_then = {"label": "the rate less the CPI's rise over the year as first reported (BLS, city wage-earner and "
                         "clerical-worker families)", "source": "treasury-bill or treasury-10-year; cpi (ALFRED CPIAUCNS, initial release)"}
    cpi_now = {"label": "the rate less the CPI-U's rise over the year today", "unit": "percent",
               "source": "FRED TB3MS or GS10 less the CPI's year-over-year change, FRED CPIAUCNS"}
    for sid, name, rate in (("real-bill-cpi", "Real bill rate, CPI", bill), ("real-10-cpi", "Real 10-year rate, CPI", ten)):
        rows = []
        for p in months:
            c, x = cpi.get(p, {}), rate[p].get("now")
            rows.append({"p": p,
                         "first": r2(x - c["yoy"]) if x is not None and c.get("yoy") is not None else None,
                         "released": c.get("released") if c.get("yoy") is not None else None,
                         "now": r2(x - c["yoy_now"]) if x is not None and c.get("yoy_now") is not None else None})
        write(series(sid, name, "M", cpi_then, cpi_now, rows))

    # quarterly, the GNP deflator
    def infl_then(q):
        qs = [q] + [prev_q(q, k) for k in (1, 2, 3)]
        ch = [defl.get(x, {}).get("chg") for x in qs]
        return None if any(v is None for v in ch) else sum(ch) / 4

    def infl_now(q):
        a, b = defl.get(q, {}).get("now"), defl.get(prev_q(q, 4), {}).get("now")
        return None if a is None or b is None else (a / b - 1) * 100

    d_then = {"label": "the quarter's average rate less the GNP deflator's rise over the year as first reported: the "
                       "mean of the four quarters' annualized changes, each as released", "source": "deflator (ALFRED GNP and GNPC96, initial release)"}
    d_now = {"label": "the quarter's average rate less the GNP deflator's rise from a year before, today", "unit": "percent",
             "source": "FRED TB3MS or GS10; FRED GNPDEF"}
    for sid, name, rate in (("real-bill", "Real bill rate, GNP deflator", bill), ("real-10", "Real 10-year rate, GNP deflator", ten)):
        rows = []
        for q in quarters:
            x = qavg(rate, q)
            it, inow = infl_then(q), infl_now(q)
            rows.append({"p": q,
                         "first": r2(x - it) if x is not None and it is not None else None,
                         "released": defl[q].get("released") if it is not None else None,
                         "now": r2(x - inow) if x is not None and inow is not None else None})
        write(series(sid, name, "Q", d_then, d_now, rows))

    # the Taylor rule
    def rule(pi, g):
        return 2 + pi + 0.5 * (pi - 2) + 0.5 * g

    rows = []
    for q in quarters:
        f = qavg(ff, q)
        it, inow = infl_then(q), infl_now(q)
        u = qavg(ue, q, "first")
        g_then = -3 * (u - 4) if u is not None else None
        g_now = gap.get(q, {}).get("now")
        rel = latest(defl[q].get("released"), ue.get(months_of(q)[-1], {}).get("released"))
        rows.append({"p": q,
                     "first": r2(f - rule(it, g_then)) if None not in (f, it, g_then) else None,
                     "released": rel if None not in (f, it, g_then) else None,
                     "now": r2(f - rule(inow, g_now)) if None not in (f, inow, g_now) else None})
    write(series("taylor-gap", "Funds rate less the Taylor rule", "Q",
                 {"label": "the quarter's average federal funds rate less Taylor's rule, with inflation as first "
                           "reported (the deflator, as above) and the gap by Okun's law from unemployment as first reported",
                  "source": "federal-funds (FRED FEDFUNDS); deflator; unemployment (ALFRED UNRATE, initial release)"},
                 {"label": "the same rule with today's deflator and the CBO's output gap", "unit": "percentage points",
                  "source": "FRED FEDFUNDS, GNPDEF; CBO potential GDP (gap-cbo)"}, rows))

    # HLW: today only
    path = os.path.join(DIR, "rstar-gap.yaml")
    if a.hlw:
        rstar, inp, final = hlw_rows(a.hlw)
        rows = []
        for q in quarters:
            if q in rstar and q in inp and None not in inp[q]:
                exp, rate = inp[q]
                rows.append({"p": q, "now": r2(rate - exp - rstar[q])})
        write(series("rstar-gap", "Real rate less the neutral rate", "Q", None,
                     {"label": "Holston, Laubach and Williams: their real policy rate (the discount rate to 1965, the "
                               "federal funds rate after, less their expected inflation, a four-quarter average of "
                               "past core consumer prices) less their natural rate, r* (US, one-sided)",
                      "unit": "percentage points",
                      "source": f"Federal Reserve Bank of New York, HLW current estimates ({final})",
                      "url": "https://www.newyorkfed.org/research/policy/rstar"}, rows))
    elif not os.path.exists(path):
        sys.exit("rstar-gap: no --hlw workbook and no file to keep")

    # the yield spread: today only
    rows = [{"p": p, "now": r2(ten[p]["now"] - bill[p]["now"])} for p in months
            if ten[p].get("now") is not None and bill[p].get("now") is not None]
    write(series("spread", "Yield spread, 10-year less the bill", "M", None,
                 {"label": "the 10-year constant-maturity yield less the 3-month bill, monthly averages",
                  "unit": "percentage points", "source": "FRED GS10 less TB3MS"}, rows))


if __name__ == "__main__":
    main()
